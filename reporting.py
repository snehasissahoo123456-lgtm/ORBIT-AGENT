"""Generate local daily ORBIT activity reports."""
from datetime import datetime, timezone
from config import settings
from memory import get_audit, get_tasks

def generate_daily_reports(day: str | None = None) -> dict:
    day=day or datetime.now(timezone.utc).date().isoformat()
    outdir=settings.workspace/"reports"; outdir.mkdir(parents=True,exist_ok=True)
    events=[e for e in get_audit(200) if str(e.get("created_at","")).startswith(day)]
    tasks=[t for t in get_tasks(100) if str(t.get("created_at","")).startswith(day)]
    pdf_path=outdir/f"orbit-daily-{day}.pdf"; ppt_path=outdir/f"orbit-daily-{day}.pptx"
    pdf_result=ppt_result=None
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        styles=getSampleStyleSheet()
        story=[Paragraph(f"ORBIT Daily Activity Report — {day}",styles["Title"]),Spacer(1,12),
               Paragraph(f"Tasks recorded: {len(tasks)} · Audit events: {len(events)}",styles["Normal"]),Spacer(1,12)]
        rows=[["Time UTC","Event","Details"]]+[[e["created_at"][11:19],e["event"],e["detail"][:150]] for e in events[:60]]
        table=Table(rows,repeatRows=1,colWidths=[50,80,320])
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#19365c")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                                   ("GRID",(0,0),(-1,-1),.25,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),7)]))
        story.append(table); SimpleDocTemplate(str(pdf_path),pagesize=letter).build(story); pdf_result=pdf_path.relative_to(settings.workspace).as_posix()
    except ImportError: pass
    try:
        from pptx import Presentation
        from pptx.util import Pt
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text=f"ORBIT Daily Report — {day}"
        slide.placeholders[1].text=f"{len(tasks)} tasks · {len(events)} audit events · Local report"
        slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text="Recent activity"
        tf=slide.placeholders[1].text_frame; tf.text="Audit events (UTC)"
        for e in events[:12]:
            p=tf.add_paragraph(); p.text=f"{e['created_at'][11:19]} — {e['event']}: {e['detail'][:100]}"; p.level=1; p.font.size=Pt(12)
        prs.save(str(ppt_path)); ppt_result=ppt_path.relative_to(settings.workspace).as_posix()
    except ImportError: pass
    return {"date":day,"tasks":len(tasks),"events":len(events),"pdf":pdf_result,"pptx":ppt_result,
            "note":"Reports summarize retained ORBIT records and cannot prove unlogged activity."}
